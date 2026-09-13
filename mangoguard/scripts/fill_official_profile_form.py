"""Fill the official-layout Student Profile Form docx.

The official form ("Corrected Student Profile Form.docx") arrived as the CREST template
with only the first name and award level filled. This script transcribes the criteria
mapping from STUDENT_PROFILE.md (the source of truth) into the form's tables.

Safe-edit workflow per crest-playbook/04: anchor every edit on existing text and fail
loudly if the anchor is missing; verify afterwards by extracting all text and checking
the expected markers. Run the verification block at the bottom after editing.

Reproduce with:  ../.venv/Scripts/python.exe scripts/fill_official_profile_form.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
FORM = ROOT / "Corrected Student Profile Form.docx"

TITLE = "Occlusion-Aware Estimation of Mango Fruit Load from Sparse Camera Viewpoints"

# (criteria row index in table 1, where-column text, notes paragraphs)
CRITERIA: list[tuple[int, str, list[str]]] = [
    (
        2,
        "§1.1–1.2, p. 2; §5.8, p. 17",
        [
            "One testable aim sentence split into five numbered objectives, each with an explicit "
            '"done when" test. Section 5.8 revisits all five against the results and states '
            "plainly which is not met: objective 5, validation against picked and counted trees, "
            "cannot be met until the fruit exists. Four are met and reported.",
        ],
    ),
    (
        3,
        "§1 opening, p. 2; §6.1, pp. 17–18",
        [
            "The grower in Ratnagiri who hired eleven pickers for a job that needed six. India "
            "grows around 22 million tonnes of mango a year and about 86% of holdings are under "
            "two hectares, so the fruit count per tree is the number everything else is planned "
            "around and almost nobody has it. Section 2 explains why the existing repair, a "
            "correction factor fitted per orchard, does not solve it.",
        ],
    ),
    (
        4,
        "§3.1, p. 6",
        [
            "Four genuinely different approaches compared in a trade-off table, each with the "
            "reason it was eliminated or chosen. The fitted multiplier was not discarded but "
            "kept as the measured baseline, because it is what published work actually does.",
        ],
    ),
    (
        5,
        "§3, pp. 6–8; Appendix B, p. 23",
        [
            "The chosen route is the only one that uses information the images already contain: "
            "a fruit seen from three viewpoints out of twelve tells you something a fruit seen "
            "from all twelve does not. Appendix B gives the seed, the population parameters and "
            "the two commands that regenerate every number and figure. The repository is public "
            "at github.com/devadit1515/mangoguard.",
        ],
    ),
    (
        6,
        "§1.3, pp. 3–4",
        [
            "Twelve stages, each ending in something checkable, with what each one depended "
            "on. The order was fixed before I started and the reason is given: the evaluation "
            "had to exist before the estimator did, so that no decision about scoring could be "
            "made by someone who already knew which answer he wanted. Two departures are "
            "stated rather than smoothed over. One stage ran later than it should have, the "
            "check of my simulated trees against a real harvested orchard, and running it late "
            "cost a regenerated set of numbers; putting it immediately after the simulator "
            "would have caught the fault in an afternoon, and that is the change I would make. "
            "The plan itself changed once, on evidence: the corrected simulator contradicted my "
            "starting premise, and I narrowed the aim to the question the measurement supported "
            "rather than adjusting the simulator until the premise survived. The three "
            "scheduled stages are timed to the fruiting season and the protocol for them is "
            "written.",
        ],
    ),
    (
        8,
        "§1.3, pp. 3–4; Appendix B, p. 23",
        [
            "Pinned tool versions (Python 3.11, NumPy 1.26, SciPy 1.11, matplotlib 3.8, pytest 8), "
            "a fixed seed, and the commands that reproduce every number. The cooperating grower "
            "in Ratnagiri is a named resource whose orchard will host the 2027 campaign. "
            "Published field measurements stood in for a mentor.",
        ],
    ),
    (
        9,
        "§2, pp. 5–6; References, pp. 21–22",
        [
            "The background is built as an argument: detection is solved (F1 0.968), the gap "
            "between seen and actual is measured (40.2% recovered), the field repairs it with a "
            "constant that varies 1.05 to 2.43, and ecology already holds the statistics for "
            "populations that hide. Fourteen references, every one cited in the text.",
        ],
    ),
    (
        11,
        "§6.1, pp. 17–18; §6.3, p. 18; §6.5, p. 19",
        [
            "Section 6.1 answers the aim with numbers, then bounds itself: the result is "
            "measured on simulated trees shown to be easier than a real hedgerow. Section 6.3 "
            "declines to set the simulated figure beside the published field figure of 11.5% "
            "and explains why the two are not comparable.",
        ],
    ),
    (
        12,
        "§6.2, p. 18; §4, pp. 9–10; §5.6, pp. 15–16",
        [
            "Replacing a smooth canopy-depth proxy with the reconstruction's measurement of "
            "unobserved volume roughly halved the error at two viewpoints. Separating the "
            "detector's hit rate from its number of chances made both identifiable. Testing "
            "against a deliberately worse detector found the one setting where the method loses "
            "to the multiplier, which is reported rather than hidden.",
        ],
    ),
    (
        13,
        "§8, p. 20",
        [
            "The simulated trees had no trunks for two days and every internal check passed, "
            "because the checks and the simulator shared the same assumptions; a published "
            "harvest count caught it. The change I would make is to go to the published field "
            "measurements before building rather than after.",
        ],
    ),
    (
        15,
        "§3.2–3.3, pp. 6–7; Appendix C, pp. 23–24",
        [
            "The body explains why re-sightings carry information about fruit never seen, and "
            "the blind spot: a fruit with no chance of being seen leaves no trace. Appendix C "
            "sets out the transmittance law, the clumping treatment, the zero-truncated "
            "likelihood and the reciprocal-probability total, each as why it is the right tool.",
        ],
    ),
    (
        16,
        "§7, pp. 19–20",
        [
            "Which error to prefer, since over-estimates idle pickers and under-estimates leave "
            "fruit past ripeness. Whose harvest is destroyed for ground truth, agreed in "
            "writing. A fairness finding measured from my own results: accuracy is worst on "
            "dense unpruned canopies, so accuracy is reported by canopy density. Dual use is "
            "named: a buyer holding the estimate while the grower does not is worse for the grower.",
        ],
    ),
    (
        17,
        "§6.4, pp. 18–19",
        [
            "Treating fruit counting as a wildlife abundance problem: the methods ecologists use "
            "for animals that hide map onto a camera walked around a tree, and I have not found "
            "the transfer made before. Using a reconstruction for what it measures rather than "
            "what it appears to measure: a map of where the cameras could not look.",
        ],
    ),
    (
        18,
        "§4.1–4.3, pp. 9–10; FIX_LOG.md in the repository",
        [
            "Three problems told in full with cause, fix and verification, a fourth that was a "
            "wrong premise rather than a code defect, and a fifth that was a wrong diagnosis I "
            "tested and refuted. Thirteen entries, five still open, in the repository's fix log.",
        ],
    ),
    (
        19,
        "Whole report; Appendix A, pp. 22–23",
        [
            "The body is written to be followed without a statistics background, with one "
            "everyday comparison per hard idea. All mathematics is quarantined in Appendix C "
            "and the argument never depends on it. Eight figures, each captioned and referred "
            "to before it appears.",
        ],
    ),
]

REFLECTIONS = [
    "Why I chose this project. A conversation my father had with the grower we buy Alphonso "
    "from every summer: he had hired eleven pickers for a harvest he thought would need six, "
    "and paid them all for a day of standing about, because his only way of estimating his "
    "crop was to walk the rows and look. I wanted to know whether that number could come "
    "from a phone instead of from experience.",
    "How it was and was not successful. It produced a simulator calibrated against published "
    "field measurements, seven estimators scored on identical trees, and a headline result: "
    "three viewpoints with the reconstruction estimator match twelve without. It did not "
    "count a single real mango. Everything in the report is measured on trees I generated, "
    "and I would rather say that plainly than let a reader assume otherwise. It also found "
    "its own limit: on two viewpoints with a poor detector the multiplier wins, and I "
    "reported that rather than hide it.",
    "What I learnt. Twice my own measurements contradicted what I had set out to show. The "
    "first time, correcting the physics of my simulator revealed that walking a full circle "
    "around a tree already recovers most of what one view hides, which undercut the premise "
    "I had started from. The second time, comparing my trees against somebody else's "
    "harvested orchard showed that mine were far too easy, because they had no trunks or "
    "branches in them at all. Both times the right move was to change the claim rather than "
    "the code, and learning to do that is worth more to me than the accuracy figure. Working "
    "without a mentor cost me on both occasions: my simulator and my checks shared the same "
    "assumptions, so everything agreed with itself and nothing caught the error. What "
    "eventually played the sceptic was published field data from an orchard someone else had "
    "harvested, which is a poor substitute for a person who is allowed to doubt you early.",
    "What impact it might have on others. The method needs no harvested trees to calibrate "
    "against, which is the practical difference: a correction factor has to be bought with a "
    "harvest, and this removes that cost. For a smallholder with a few hundred trees, the "
    "difference between three viewpoints and twelve is the difference between scanning an "
    "orchard in a morning and spending a whole day at it. Whether any of this holds on real "
    "trees is a harvest away.",
    "What I would improve. I would have gone to the published field measurements before "
    "building the simulator rather than after. The 40.2% figure was available the whole time "
    "and would have caught the missing wood on the first day. I built inward from physics I "
    "trusted instead of outward from measurements someone had already made. Above all I "
    "would find one person allowed to doubt me early, because explaining a result to someone "
    "else catches what re-reading my own code never did.",
]

NOTES_BOX = (
    "Status. The project is at its simulation stage. The estimator, the simulator that tests "
    "it and the comparison against current practice are complete, and every number comes from "
    "one seeded script. The orchard validation runs between February and May 2027, because "
    "the ground truth for a tree is obtained by picking it and counting, and the fruit does "
    "not exist before then.\n\n"
    "Reproducing the work. Everything in the report can be re-run: two commands regenerate "
    "every number and redraw every figure from a fixed seed, and a test suite of twenty "
    "tests checks the simulator against closed-form answers rather than against itself. The "
    "repository is public at github.com/devadit1515/mangoguard.\n\n"
    'How I used AI. Declared in full in the report under "A note on AI use". In short, I '
    "used Anthropic's Claude substantially: it wrote most of the code from my direction and "
    "drafted sections of the report from my decisions about what it should argue. I set the "
    "aim and the success conditions before any code existed, I decided to narrow the aim "
    "rather than adjust the simulator when the corrected physics contradicted it, and I can "
    "explain the method, the statistics and the limitations of this work. The commit history "
    "holds the dated trail."
)


def _clear_cell(cell) -> None:
    """Remove every paragraph but the first, and empty the first."""
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first = cell.paragraphs[0]
    for run in list(first.runs):
        run._element.getparent().remove(run._element)


def _write_cell(cell, text: str, bold: bool = False) -> None:
    _clear_cell(cell)
    first = True
    for para_text in text.split("\n\n"):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = p.add_run(para_text)
        run.bold = bold


def fill(doc: Document) -> None:
    # Header table: project title.
    header = doc.tables[0]
    assert "Project title" in header.rows[2].cells[0].text
    _write_cell(header.rows[2].cells[1], TITLE)

    # Criteria table: Where column and Notes column.
    criteria = doc.tables[1]
    example = criteria.rows[2].cells[1].text
    assert "Example" in example or "§1.1" in example, "unexpected content in row 2 Where cell"
    for row_idx, where, notes in CRITERIA:
        cells = criteria.rows[row_idx].cells
        _write_cell(cells[1], where)
        _write_cell(cells[2], "\n\n".join(notes))

    # Reflections box: replace the prompt bullets.
    tail = doc.tables[2]
    refl = tail.rows[1].cells[0].text
    assert (
        "How my project was successful" in refl or "Why I chose this project" in refl
    ), "unexpected content in reflections box"
    _write_cell(tail.rows[1].cells[0], "\n\n".join(REFLECTIONS))

    # Further-notes box.
    notes = tail.rows[7].cells[0].text
    assert (
        "further notes" in notes.lower() or "Status." in notes
    ), "unexpected content in further-notes box"
    _write_cell(tail.rows[7].cells[0], NOTES_BOX)


def _ascii(s: str) -> str:
    return s.encode("ascii", "backslashreplace").decode("ascii")


def verify(path: Path) -> None:
    doc = Document(path)
    all_text = []
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                all_text.append(c.text)
    text = "\n".join(all_text)
    expected = [
        TITLE,
        "irteen entries",
        "github.com/devadit1515/mangoguard",
        "40.2%",
        "five still open",
        "harvest away",
        "change the claim rather than the code",
    ]
    lowered = text.lower()
    missing = [m for m in expected if m.lower() not in lowered]
    assert not missing, f"missing markers: {missing}"
    assert "Example: Page 2" not in text, "template example row still present"
    empty_where = 0
    for row_idx, _where, _ in CRITERIA:
        cells = doc.tables[1].rows[row_idx].cells
        if not cells[1].text.strip() or not cells[2].text.strip():
            empty_where += 1
    assert empty_where == 0, f"{empty_where} criteria rows still empty"
    print(f"verified {path.name}: title, 15 criteria rows, reflections and notes all filled")


if __name__ == "__main__":
    document = Document(FORM)
    fill(document)
    document.save(FORM)
    verify(FORM)
